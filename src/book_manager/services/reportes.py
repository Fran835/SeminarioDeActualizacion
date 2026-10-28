"""Generación del reporte Excel de comparación de precios."""

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from book_manager.services.auditoria import auditar
from book_manager.services.comparacion import ServicioComparacion

CAMPOS = (
  "Libro",
  "Precio interno",
  "Precio web",
  "Diferencia",
  "Fecha de extracción",
)


class ServicioReportes:
  """Genera el reporte Excel con las comparaciones de precios.

  Args:
    comparacion: Servicio que obtiene las comparaciones.
    carpeta_reportes: Carpeta de destino; por defecto, la de comparación.
  """

  def __init__(
    self,
    comparacion: ServicioComparacion,
    carpeta_reportes: Optional[Path] = None,
  ) -> None:
    self._comparacion = comparacion
    self._carpeta = Path(
      carpeta_reportes or comparacion.carpeta_reportes
    )

  @property
  def carpeta_reportes(self) -> Path:
    """Path: Carpeta donde se guardan los reportes Excel."""
    return self._carpeta

  @auditar("generar_reporte")
  def generar_reporte_precios(
    self, ruta_jsonl: Optional[Path] = None
  ) -> Path:
    """Genera un Excel con las comparaciones más recientes.

    Args:
      ruta_jsonl: JSONL a procesar; si es None, se usa el más reciente.

    Returns:
      Path: Ruta del archivo Excel generado.

    Raises:
      ValueError: Si no hay comparaciones para exportar.
    """
    comparaciones = self._comparacion.comparar(ruta_jsonl)
    if not comparaciones:
      raise ValueError(
        "No hay comparaciones para generar el reporte. "
        "Ejecutá el scraping y verificá que los títulos coincidan."
      )

    self._carpeta.mkdir(parents=True, exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    ruta = self._carpeta / f"reporte_precios_{marca}.xlsx"

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Precios"
    hoja.append(CAMPOS)
    for celda in hoja[1]:
      celda.font = Font(bold=True, color="FFFFFF")
      celda.fill = PatternFill("solid", fgColor="1F4E78")
      celda.alignment = Alignment(horizontal="center")

    for comparacion in comparaciones:
      hoja.append([
        comparacion.libro,
        float(comparacion.precio_interno),
        float(comparacion.precio_web),
        float(comparacion.diferencia),
        comparacion.fecha_extraccion,
      ])

    for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=4):
      for celda in fila:
        celda.number_format = '#,##0.00;[Red]-#,##0.00'
        celda.alignment = Alignment(horizontal="right")
    for celda in hoja["E"][1:]:
      celda.number_format = "yyyy-mm-dd hh:mm:ss"
      celda.alignment = Alignment(horizontal="center")

    anchos = (38, 16, 16, 16, 22)
    for indice, ancho in enumerate(anchos, start=1):
      hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:E{hoja.max_row}"

    libro.save(ruta)
    return ruta
